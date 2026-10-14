import pptx
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PPT_PATH = r"C:\Users\RONAK AGRAWAL\OneDrive\Desktop\Smart_Face_Attendance_System.pptx"

def create_operational_slides():
    prs = pptx.Presentation(PPT_PATH)
    blank_layout = prs.slide_layouts[6] # Blank layout

    # Constants matching existing slides
    BANNER_COLOR = RGBColor(0x00, 0x33, 0x66)       # Dark Navy
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    TEXT_DARK = RGBColor(0x28, 0x28, 0x28)          # Charcoal dark
    SUBTITLE_COLOR = RGBColor(0x47, 0x55, 0x69)     # Slate
    ACCENT_BLUE = RGBColor(0x02, 0x84, 0xC7)        # Sky/Blue Accent
    HIGHLIGHT_BG = RGBColor(0xF1, 0xF5, 0xF9)       # Soft grey card bg
    CARD_BORDER = RGBColor(0xCB, 0xD5, 0xE1)

    slides_data = [
        {
            "badge": "11",
            "title": "System Execution & Startup Guide",
            "subtitle": "1-Click Automation & Manual Launch Commands",
            "sections": [
                {
                    "heading": "1-Click Master Startup (Recommended)",
                    "bullets": [
                        "Simply double-click the included master script: start_all.bat in the project root.",
                        "Location: C:\\Users\\RONAK AGRAWAL\\Projects\\smart-attendance-system\\start_all.bat",
                        "Automatically launches Central FastAPI Backend + all 4 Next.js Web Portals in separate consoles."
                    ]
                },
                {
                    "heading": "Manual Component-by-Component Startup Commands",
                    "bullets": [
                        "1. Backend API (FastAPI): cd backend && venv\\Scripts\\python.exe -m uvicorn main:app --reload --port 8000",
                        "2. Web Showcase & Teacher Console: cd web-frontend && npm run dev -- -p 3000",
                        "3. Admin Governance Portal: cd admin-portal && npm run dev -- -p 3001",
                        "4. Student Academic Portal: cd student-portal && npm run dev -- -p 3002",
                        "5. HOD Departmental Command Hub: cd hod-portal && npm run dev -- -p 3003",
                        "6. Teacher Mobile App (Expo): cd mobile-app && npx expo start (or scan QR in Expo Go)"
                    ]
                },
                {
                    "heading": "1-Click Graceful Shutdown",
                    "bullets": [
                        "Double-click stop_all.bat in the project root to cleanly terminate all processes and release ports 8000, 3000-3003."
                    ]
                }
            ]
        },
        {
            "badge": "12",
            "title": "Live Service Endpoints & Network Topology",
            "subtitle": "Complete Access URLs, Ports, and Swagger Documentation",
            "sections": [
                {
                    "heading": "Backend API & Interactive Documentation",
                    "bullets": [
                        "REST API Root Engine: http://localhost:8000",
                        "Swagger Interactive Documentation: http://localhost:8000/docs (Live API testing & schema explorer)",
                        "ReDoc OpenAPI Specification: http://localhost:8000/redoc"
                    ]
                },
                {
                    "heading": "Web Application Portals & Dashboards",
                    "bullets": [
                        "Admin Governance Portal: http://localhost:3001 (Master Settings, Defaulters, Student/Faculty Directory)",
                        "HOD Command Deck: http://localhost:3003 (Classes, Faculty Oversight, Bunking Alerts, Excel Reports)",
                        "Student Academic Portal: http://localhost:3002 (Live Attendance Tracking, Selfie Enrollment, Grievance Filing)",
                        "Classroom Teacher Web Terminal: http://localhost:3000/teacher (Dynamic Class Roster, Instant Attendance Commit)"
                    ]
                },
                {
                    "heading": "Mobile Biometric Terminal (Expo / React Native)",
                    "bullets": [
                        "Host URL: http://<LAN_IP>:8000 or Cloud Tunnel URL configured in mobile-app/App.js",
                        "Features: Apple Dynamic Island live animations, Haptic feedback, Walk-up Kiosk Mode, Offline Queue sync."
                    ]
                }
            ]
        },
        {
            "badge": "13",
            "title": "System Credentials & Authentication Vault",
            "subtitle": "Active User Accounts, Passwords, Role Scopes & Login Access",
            "sections": [
                {
                    "heading": "1. Super Administrator (Master Governance Account)",
                    "bullets": [
                        "Portal URL: http://localhost:3001",
                        "Email / Username: ronakagrawalmthr@gmail.com",
                        "Admin Name: Victor Von Doom (Institutional Administrator, Victus University)",
                        "Password: Password set during System Setup (Can be reset anytime via real Email OTP verification)",
                        "Access Scope: University Configuration, Master CRUD, Defaulter Thresholds, System-wide Audits."
                    ]
                },
                {
                    "heading": "2. Student Portal Account (Enrolled Student Profile)",
                    "bullets": [
                        "Portal URL: http://localhost:3002",
                        "Student Roll Number: 2500650100111  |  Email: 2500650100111@student.university.edu",
                        "Password: 2500650100111",
                        "Profile: Victor Von Doom | Branch: Computer Science & Engineering | Semester: 6 | Section: A",
                        "Access Scope: Real-time subject-wise percentage, attendance history timeline, biometric status."
                    ]
                },
                {
                    "heading": "3. Faculty / Teacher & HOD Accounts",
                    "bullets": [
                        "Teacher Console: Mobile App or http://localhost:3000/teacher (Login via faculty email + password)",
                        "HOD Hub: http://localhost:3003 (Department-scoped monitoring, holiday declarations, analytics)"
                    ]
                }
            ]
        },
        {
            "badge": "14",
            "title": "End-to-End Operational Workflow Guide",
            "subtitle": "Step-by-Step Instructions for Daily Classroom Attendance & Governance",
            "sections": [
                {
                    "heading": "Phase 1: Classroom Setup & Session Selection",
                    "bullets": [
                        "Teacher logs in via Mobile App or Web Terminal (http://localhost:3000/teacher).",
                        "Selects Branch (e.g. Computer Science), Semester (6), Section (A), and Subject (e.g. CS305 - DAA).",
                        "Student roster loads dynamically from SQLite database (all students default to unmarked/absent)."
                    ]
                },
                {
                    "heading": "Phase 2: High-Speed Biometric Verification & Liveness Check",
                    "bullets": [
                        "Kiosk mode is launched: Students walk up to camera one by one.",
                        "Anti-spoofing engine verifies real-time human liveness (prevents digital screen & printed photo attacks).",
                        "Sub-second face matching matches 128-d vector against enrolled class roster.",
                        "Dynamic Island pops up: Present confirmed with green badge and tactile haptic vibration."
                    ]
                },
                {
                    "heading": "Phase 3: Review, Manual Overrides & Final Record Commit",
                    "bullets": [
                        "Teacher reviews present and absent counts; applies manual overrides for medical/injury exceptions.",
                        "Enters optional notes or Guest Conductor name, then clicks Final Submit.",
                        "Database commits record with strict anti-duplicate protection; HOD & Student dashboards update instantly."
                    ]
                }
            ]
        },
        {
            "badge": "15",
            "title": "Production Hardening, Security & Maintenance",
            "subtitle": "Zero-Backdoor Security, Concurrency Optimization & Data Safeguards",
            "sections": [
                {
                    "heading": "Zero-Backdoor Security Architecture",
                    "bullets": [
                        "All backdoor master OTPs ('123456') and unauthenticated password bypasses permanently eliminated.",
                        "Forgot Password & Setup OTPs dispatched securely via verified institutional email service (Resend).",
                        "Full JWT Bearer authorization enforced on all attendance submission and management endpoints."
                    ]
                },
                {
                    "heading": "Database Concurrency & Integrity Hardening",
                    "bullets": [
                        "SQLite WAL (Write-Ahead Logging) enabled with 10,000ms busy timeout to prevent multi-portal write locks.",
                        "Database-level UniqueConstraint(session_id, student_id) guarantees zero duplicate attendance records.",
                        "N+1 query bottleneck solved: Defaulters analytics executes via a single optimized SQL GROUP BY query."
                    ]
                },
                {
                    "heading": "Automated Midnight Cron & Clean Shutdown",
                    "bullets": [
                        "Midnight Cron Worker calculates exact seconds until 00:00 AM for automated daily archival & notifications.",
                        "Database session leaks prevented via explicit try...finally cleanup in all background jobs."
                    ]
                }
            ]
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # 1. Header Banner Shape (Rectangle 1)
        top_banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
        top_banner.fill.solid()
        top_banner.fill.fore_color.rgb = BANNER_COLOR
        top_banner.line.color.rgb = BANNER_COLOR

        # 2. Slide Number Badge Box (TextBox 2)
        badge_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(1.0), Inches(0.8))
        tf_badge = badge_box.text_frame
        tf_badge.word_wrap = True
        p_badge = tf_badge.paragraphs[0]
        p_badge.text = data["badge"]
        p_badge.font.size = Pt(28)
        p_badge.font.bold = True
        p_badge.font.color.rgb = WHITE
        p_badge.alignment = PP_ALIGN.LEFT

        # 3. Slide Title (TextBox 3)
        title_box = slide.shapes.add_textbox(Inches(1.5), Inches(0.25), Inches(11.0), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        # 4. Slide Subtitle (TextBox 4)
        sub_box = slide.shapes.add_textbox(Inches(1.5), Inches(1.35), Inches(11.0), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(14)
        p_sub.font.bold = True
        p_sub.font.color.rgb = SUBTITLE_COLOR

        # 5. Main Content Box (TextBox 5)
        content_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.85), Inches(12.1), Inches(4.7))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True

        for s_idx, sec in enumerate(data["sections"]):
            # Section Heading
            p_head = tf_content.add_paragraph() if (s_idx > 0 or len(tf_content.paragraphs[0].text) > 0) else tf_content.paragraphs[0]
            p_head.text = "◆  " + sec["heading"]
            p_head.font.size = Pt(16)
            p_head.font.bold = True
            p_head.font.color.rgb = BANNER_COLOR
            p_head.space_before = Pt(8) if s_idx > 0 else Pt(0)
            p_head.space_after = Pt(3)

            # Bullet points
            for bullet in sec["bullets"]:
                p_bullet = tf_content.add_paragraph()
                p_bullet.text = "    •  " + bullet
                p_bullet.font.size = Pt(13)
                p_bullet.font.color.rgb = TEXT_DARK
                p_bullet.space_before = Pt(1)
                p_bullet.space_after = Pt(2)

    prs.save(PPT_PATH)
    print(f"Successfully added {len(slides_data)} operational slides to {PPT_PATH}!")

if __name__ == "__main__":
    create_operational_slides()
