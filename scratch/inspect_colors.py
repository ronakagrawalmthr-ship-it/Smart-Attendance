import pptx

prs = pptx.Presentation(r'C:\Users\RONAK AGRAWAL\OneDrive\Desktop\Smart_Face_Attendance_System.pptx')
s = prs.slides[1]
print("Total shapes in slide 2:", len(s.shapes))
for shp in s.shapes:
    print(f"Shape: {shp.name}, shape_type: {shp.shape_type}")
    try:
        if shp.fill.type is not None:
            print(f"  fill type: {shp.fill.type}")
            try:
                print(f"  fill color: {shp.fill.fore_color.rgb}")
            except Exception as e:
                print(f"  fill color error: {e}")
    except Exception:
        pass
