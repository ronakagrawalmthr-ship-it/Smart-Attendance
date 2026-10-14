import pptx
from pptx.enum.dml import MSO_COLOR_TYPE

prs = pptx.Presentation(r'C:\Users\RONAK AGRAWAL\OneDrive\Desktop\Smart_Face_Attendance_System.pptx')
print(f'Width: {prs.slide_width}, Height: {prs.slide_height}')
s1 = prs.slides[0]
print("--- Slide 1 Shapes ---")
for shape in s1.shapes:
    print(f'Shape: {shape.name}, Type: {shape.shape_type}')
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                c = r.font.color.rgb if (r.font.color and r.font.color.type == MSO_COLOR_TYPE.RGB) else 'None'
                print(f'   Run: "{r.text[:35]}", Font: {r.font.name}, Size: {r.font.size}, Color: {c}')

s2 = prs.slides[1]
print("--- Slide 2 Shapes ---")
for shape in s2.shapes:
    print(f'Shape: {shape.name}, Type: {shape.shape_type}')
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                c = r.font.color.rgb if (r.font.color and r.font.color.type == MSO_COLOR_TYPE.RGB) else 'None'
                print(f'   Run: "{r.text[:35]}", Font: {r.font.name}, Size: {r.font.size}, Color: {c}')
